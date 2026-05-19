from datetime import datetime
from pathlib import Path

from app.services.pass_db import CI, PASS_TYPE_TEMPORARY, fetch_pass_by_qr, pass_status
from app.services.qr_codes import save_qr_code
from app.services.reference_data import remember_reference_values
from app.services.validation import ValidationError, validate_pass_data


TEMPORARY_TYPE = PASS_TYPE_TEMPORARY
TEMP_STATUS_FREE = "free"
TEMP_STATUS_ISSUED = "issued"
TEMP_STATUS_EXPIRED = "expired"
TEMP_STATUS_LOST = "lost"
TEMP_POOL_PREFIX = "TMP"
TEMP_POOL_SIZE = 500

TEMP_STATUS_TITLES = {
    TEMP_STATUS_FREE: "Свободен",
    TEMP_STATUS_ISSUED: "Выдан",
    TEMP_STATUS_EXPIRED: "Истёк",
    TEMP_STATUS_LOST: "Утерян",
}


class TemporaryPassError(ValueError):
    pass


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def temporary_qr(index, prefix=TEMP_POOL_PREFIX, book_no=None):
    if book_no is None:
        return f"{prefix}-{index:04d}"
    return f"{prefix}-B{int(book_no):03d}-{index:04d}"


def is_temporary_pass(row):
    return bool(row and len(row) > CI["type"] and row[CI["type"]] == TEMPORARY_TYPE)


def temporary_status(row):
    if not is_temporary_pass(row):
        return ""
    return row[CI["temp_status"]] or TEMP_STATUS_FREE


def temporary_book_no(row):
    if not is_temporary_pass(row) or len(row) <= CI["temp_book"]:
        return 0
    return int(row[CI["temp_book"]] or 0)


def temporary_book_title(book_no):
    return f"Книга №{book_no}" if book_no else "Без книги"


def temporary_status_title(status):
    return TEMP_STATUS_TITLES.get(status, status or "")


def _name(row):
    return " ".join(filter(None, [row[CI["ln"]], row[CI["fn"]], row[CI["mn"]]]))


def _book_size(book):
    return int(book[3] or TEMP_POOL_SIZE) if book else TEMP_POOL_SIZE


def _next_book_no(db):
    cursor = db.cursor()
    row = cursor.execute("SELECT COALESCE(MAX(book_no), 0) + 1 FROM temporary_books").fetchone()
    return int(row[0] or 1)


def _ensure_book_record(db, book_no, size):
    db.execute(
        "INSERT OR IGNORE INTO temporary_books (book_no,created_at,size) VALUES (?,?,?)",
        (book_no, _now(), int(size or TEMP_POOL_SIZE)),
    )
    db.execute(
        "UPDATE temporary_books SET size=MAX(COALESCE(size, 0), ?) WHERE book_no=?",
        (int(size or TEMP_POOL_SIZE), book_no),
    )


def active_temporary_book(db):
    cursor = db.cursor()
    return cursor.execute(
        """SELECT book_no, created_at, completed_at, size
           FROM temporary_books
           WHERE completed_at IS NULL
           ORDER BY book_no DESC
           LIMIT 1"""
    ).fetchone()


def latest_unarchived_completed_book(db):
    cursor = db.cursor()
    return cursor.execute(
        """SELECT book_no, created_at, completed_at, size, archived_at, archive_path
           FROM temporary_books
           WHERE completed_at IS NOT NULL
             AND (archived_at IS NULL OR archived_at='')
           ORDER BY book_no ASC
           LIMIT 1"""
    ).fetchone()


def mark_temporary_book_archived(db, book_no, archive_path):
    db.execute(
        "UPDATE temporary_books SET archived_at=?, archive_path=? WHERE book_no=?",
        (_now(), str(archive_path or ""), int(book_no)),
    )
    db.commit()


def export_temporary_book_xlsx(db, book_no, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = db.execute(
        "SELECT book_no, created_at, completed_at, size FROM temporary_books WHERE book_no=?",
        (int(book_no),),
    ).fetchone()
    if not book:
        raise TemporaryPassError("QR-книга не найдена")

    rows = list_temporary_passes(db, "", "all", book_no=int(book_no))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Книга {int(book_no)}"

    ws["A1"] = temporary_book_title(book_no)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Создана: {book[1] or ''}"
    ws["A3"] = f"Закрыта: {book[2] or ''}"
    ws["A4"] = f"Размер: {book[3] or TEMP_POOL_SIZE}"

    headers = ["№", "QR", "Статус", "ФИО", "Куда", "Основание", "Выдан", "Срок, сут.", "Телефон"]
    header_row = 6
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6FEB")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=header_row + 1):
        values = [
            row[CI["temp_number"]] or row_index - header_row,
            row[CI["qr"]],
            temporary_status_title(temporary_status(row)),
            _name(row),
            row[CI["unit"]] or row[CI["district"]] or "",
            row[CI["rank"]] or "",
            row[CI["issued"]] or "",
            row[CI["days"]] or "",
            row[CI["phone"]] or "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col, value=value)

    widths = [8, 20, 14, 30, 26, 28, 14, 12, 18]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=col).column_letter].width = width
    ws.freeze_panes = "A7"
    wb.save(path)
    return path


def _book_total_count(db, book_no):
    return db.execute(
        "SELECT COUNT(*) FROM passes WHERE pass_type=? AND temp_book=?",
        (TEMPORARY_TYPE, book_no),
    ).fetchone()[0]


def _book_free_count(db, book_no):
    return db.execute(
        "SELECT COUNT(*) FROM passes WHERE pass_type=? AND temp_book=? AND deleted=0 AND temp_status=?",
        (TEMPORARY_TYPE, book_no, TEMP_STATUS_FREE),
    ).fetchone()[0]


def _book_uses_legacy_qr_format(db, book_no):
    return db.execute(
        """SELECT COUNT(*) FROM passes
           WHERE pass_type=? AND temp_book=? AND qr_code LIKE ? AND qr_code NOT LIKE ?""",
        (TEMPORARY_TYPE, book_no, f"{TEMP_POOL_PREFIX}-%", f"{TEMP_POOL_PREFIX}-B%-%"),
    ).fetchone()[0] > 0


def _qr_for_book(db, index, prefix, book_no):
    if _book_uses_legacy_qr_format(db, book_no):
        return temporary_qr(index, prefix)
    return temporary_qr(index, prefix, book_no)


def complete_temporary_book_if_used(db, book_no=None):
    book = active_temporary_book(db) if book_no is None else db.execute(
        "SELECT book_no, created_at, completed_at, size FROM temporary_books WHERE book_no=?",
        (book_no,),
    ).fetchone()
    if not book or book[2]:
        return None

    current_book_no = int(book[0])
    size = _book_size(book)
    if _book_total_count(db, current_book_no) >= size and _book_free_count(db, current_book_no) == 0:
        db.execute(
            "UPDATE temporary_books SET completed_at=? WHERE book_no=? AND completed_at IS NULL",
            (_now(), current_book_no),
        )
        db.commit()
        return {"book_no": current_book_no, "size": size}
    return None


def create_temporary_pool(db, count=TEMP_POOL_SIZE, prefix=TEMP_POOL_PREFIX, create_qr_files=True, book_no=None):
    complete_temporary_book_if_used(db)
    book = None
    if book_no is None:
        book = active_temporary_book(db)
        if book:
            book_no = int(book[0])
            count = _book_size(book)
        else:
            book_no = _next_book_no(db)
    else:
        book_no = int(book_no)

    _ensure_book_record(db, book_no, count)
    created = []
    for index in range(1, int(count) + 1):
        qr_code = _qr_for_book(db, index, prefix, book_no)
        cursor = db.execute(
            """INSERT OR IGNORE INTO passes
               (qr_code,district,unit,rank,last_name,first_name,middle_name,phone,
                issued_date,days_count,photo_path,active,deleted,pass_type,temp_status,temp_book,temp_number)
               VALUES (?, '', '', '', '', '', '', '', '', 1, NULL, 0, 0, ?, ?, ?, ?)""",
            (qr_code, TEMPORARY_TYPE, TEMP_STATUS_FREE, book_no, index),
        )
        if cursor.rowcount:
            created.append(qr_code)
            if create_qr_files:
                save_qr_code(qr_code)
    db.commit()
    return created


def refresh_temporary_statuses(db):
    cursor = db.cursor()
    cursor.execute(
        "SELECT * FROM passes WHERE pass_type=? AND temp_status=? AND deleted=0",
        (TEMPORARY_TYPE, TEMP_STATUS_ISSUED),
    )
    expired = []
    for row in cursor.fetchall():
        active, _left, _expires = pass_status(db, row)
        if not active:
            expired.append(row[CI["qr"]])
    if expired:
        db.executemany(
            "UPDATE passes SET active=0,temp_status=? WHERE qr_code=?",
            [(TEMP_STATUS_EXPIRED, qr_code) for qr_code in expired],
        )
        db.commit()
    complete_temporary_book_if_used(db)
    return expired


def temporary_counts(db, book_no=None):
    refresh_temporary_statuses(db)
    cursor = db.cursor()
    counts = {status: 0 for status in TEMP_STATUS_TITLES}
    params = [TEMPORARY_TYPE]
    where = "pass_type=?"
    if book_no is not None:
        where += " AND temp_book=?"
        params.append(int(book_no))
    cursor.execute(f"SELECT temp_status, COUNT(*) FROM passes WHERE {where} GROUP BY temp_status", params)
    for status, count in cursor.fetchall():
        counts[status or TEMP_STATUS_FREE] = count
    counts["total"] = sum(counts.values())
    book = active_temporary_book(db)
    counts["active_book"] = int(book[0]) if book else None
    counts["active_book_size"] = _book_size(book) if book else TEMP_POOL_SIZE
    counts["active_book_free"] = _book_free_count(db, int(book[0])) if book else 0
    return counts


def list_temporary_passes(db, search="", status="all", book_no=None):
    refresh_temporary_statuses(db)
    clauses = ["pass_type=?", "deleted=0"]
    params = [TEMPORARY_TYPE]
    if book_no is not None:
        clauses.append("temp_book=?")
        params.append(int(book_no))
    if status and status != "all":
        clauses.append("temp_status=?")
        params.append(status)
    if search:
        searchable = ("qr_code", "district", "unit", "rank", "last_name", "first_name", "middle_name", "phone", "issued_date")
        for term in str(search).split():
            like = f"%{term}%"
            clauses.append("(" + " OR ".join(f"{column} LIKE ?" for column in searchable) + ")")
            params.extend([like] * len(searchable))
    cursor = db.cursor()
    cursor.execute(
        f"""SELECT * FROM passes
            WHERE {' AND '.join(clauses)}
            ORDER BY temp_book DESC, temp_number ASC, qr_code ASC""",
        params,
    )
    return cursor.fetchall()


def next_free_temporary_pass(db, create_qr_files=True):
    refresh_temporary_statuses(db)
    create_temporary_pool(db, create_qr_files=create_qr_files)
    book = active_temporary_book(db)
    if not book:
        return None
    cursor = db.cursor()
    cursor.execute(
        """SELECT * FROM passes
           WHERE pass_type=? AND temp_status=? AND deleted=0 AND temp_book=?
           ORDER BY temp_number ASC, qr_code ASC
           LIMIT 1""",
        (TEMPORARY_TYPE, TEMP_STATUS_FREE, int(book[0])),
    )
    row = cursor.fetchone()
    if row:
        return row

    complete_temporary_book_if_used(db, int(book[0]))
    create_temporary_pool(db, create_qr_files=create_qr_files)
    book = active_temporary_book(db)
    if not book:
        return None
    cursor.execute(
        """SELECT * FROM passes
           WHERE pass_type=? AND temp_status=? AND deleted=0 AND temp_book=?
           ORDER BY temp_number ASC, qr_code ASC
           LIMIT 1""",
        (TEMPORARY_TYPE, TEMP_STATUS_FREE, int(book[0])),
    )
    return cursor.fetchone()


def _require_temporary(db, qr_code):
    row = fetch_pass_by_qr(db, qr_code, include_deleted=True)
    if not row or not is_temporary_pass(row):
        raise TemporaryPassError("Временный QR не найден")
    return row


def issue_temporary_pass(db, qr_code, data, create_next_book=True, create_qr_files=True):
    row = _require_temporary(db, qr_code)
    if temporary_status(row) != TEMP_STATUS_FREE:
        raise TemporaryPassError("Этот одноразовый QR уже использован и не может быть выдан повторно")

    payload = dict(data)
    payload["qr_code"] = qr_code
    payload["photo_path"] = None
    payload["active"] = 1
    cleaned = validate_pass_data(payload, db, existing_qr=qr_code)
    if cleaned["days_count"] < 1 or cleaned["days_count"] > 10:
        raise ValidationError(["Срок временного пропуска должен быть от 1 до 10 суток"])

    db.execute(
        """UPDATE passes SET district=?,unit=?,rank=?,last_name=?,first_name=?,middle_name=?,
           phone=?,issued_date=?,days_count=?,photo_path=NULL,active=1,deleted=0,
           pass_type=?,temp_status=? WHERE qr_code=?""",
        (
            cleaned["district"],
            cleaned["unit"],
            cleaned["rank"],
            cleaned["last_name"],
            cleaned["first_name"],
            cleaned["middle_name"],
            cleaned["phone"],
            cleaned["issued_date"],
            cleaned["days_count"],
            TEMPORARY_TYPE,
            TEMP_STATUS_ISSUED,
            qr_code,
        ),
    )
    db.commit()
    remember_reference_values(db, cleaned)
    updated = fetch_pass_by_qr(db, qr_code, include_deleted=True)
    completed = complete_temporary_book_if_used(db, temporary_book_no(updated))
    if completed and create_next_book:
        create_temporary_pool(db, count=completed["size"], create_qr_files=create_qr_files)
    return updated


def return_temporary_pass(db, qr_code):
    _require_temporary(db, qr_code)
    raise TemporaryPassError("Одноразовый QR нельзя вернуть в свободные: книга сохраняет историю выдачи")


def mark_temporary_lost(db, qr_code, create_next_book=True, create_qr_files=True):
    row = _require_temporary(db, qr_code)
    db.execute(
        "UPDATE passes SET active=0,deleted=0,pass_type=?,temp_status=? WHERE qr_code=?",
        (TEMPORARY_TYPE, TEMP_STATUS_LOST, qr_code),
    )
    db.commit()
    updated = fetch_pass_by_qr(db, qr_code, include_deleted=True)
    completed = complete_temporary_book_if_used(db, temporary_book_no(updated))
    if completed and create_next_book:
        create_temporary_pool(db, count=completed["size"], create_qr_files=create_qr_files)
    return {"qr_code": qr_code, "name": _name(row), "status": temporary_status(updated), "lost_at": _now()}
