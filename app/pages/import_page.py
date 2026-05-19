import os
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

from app.core.config import C, FONT as F
from app.core.logging import get_logger
from app.services.audit import safe_record_action
from app.services.auth import PERMISSION_IMPORT, has_permission
from app.services.kpp_exchange import (PACKAGE_EXTENSION, export_kpp_package,
                                       import_kpp_package,
                                       inspect_kpp_package)
from app.services.pass_db import delete_pass_forever, insert_pass_ignore
from app.services.qr_codes import save_qr_code
from app.services.validation import ValidationError, validate_pass_data
from app.ui.widgets import Btn, _sep

logger = get_logger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    HAS_XL = True
except ImportError:
    HAS_XL = False


def _desktop_dir():
    desktop = Path.home() / "Desktop"
    return desktop if desktop.exists() else Path.home()


def _package_summary(info):
    counts = info.get("counts") or {}
    monthly = counts.get("regular", 0)
    semiannual = counts.get("semiannual", 0)
    return f"Месячные: {monthly}\nПолугодовые: {semiannual}\nВсего: {info.get('count', monthly + semiannual)}"

def show_import(app):
    if not has_permission(app.user, PERMISSION_IMPORT):
        app._toast("Недостаточно прав")
        return
    if not HAS_XL:
        messagebox.showerror("Ошибка","Установите openpyxl:\npip install openpyxl"); return

    win=app._modal("📥  Импорт из Excel",780,720, scroll=True)

    # Инструкция
    inf=tk.Frame(win,bg=C["input"],highlightthickness=1,highlightbackground=C["border"])
    inf.pack(fill="x",padx=24,pady=(12,16))
    tk.Label(inf,bg=C["input"],fg=C["mid"],font=(F,10),justify="left",
             text="Ожидаемые колонки (первая строка — заголовок, пропускается):\n"
                  "Округ  |  В/ч  |  Звание  |  Фамилия  |  Имя  |  Отчество  |  "
                  "Номер телефона  |  Дата получения  |  Количество суток\n"
                  "Дата: ДД.ММ.ГГГГ или ГГГГ-ММ-ДД.  Суток: 1–10 или 30 (месяц).").pack(padx=12,pady=10)

    # Путь
    tk.Label(win,text="Файл Excel",bg=C["panel"],fg=C["muted"],font=(F,9)).pack(anchor="w",padx=28,pady=(0,4))
    pf=tk.Frame(win,bg=C["panel"]); pf.pack(fill="x",padx=28)
    pv=tk.StringVar()
    pe=tk.Entry(pf,bg=C["input"],fg=C["text"],insertbackground=C["accent"],
                textvariable=pv,relief="flat",font=(F,11),
                highlightthickness=1,highlightbackground=C["border"])
    pe.pack(side="left",fill="x",expand=True,ipady=7,padx=(0,8))
    def browse():
        f=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx")])
        if f: pv.set(f)
    Btn(pf,text="Обзор",cmd=browse,variant="ghost",w=80,h=36,fs=10,bg=C["panel"]).pack(side="left")

    _sep(win).pack(fill="x",padx=24,pady=12)

    tk.Label(win,text="Журнал импорта",bg=C["panel"],fg=C["muted"],font=(F,9)).pack(anchor="w",padx=28,pady=(0,4))
    logf=tk.Frame(win,bg=C["input"]); logf.pack(fill="both",expand=True,padx=28)
    logt=tk.Text(logf,bg=C["input"],fg=C["text"],font=(F,10),
                 relief="flat",state="disabled",height=7,wrap="word",highlightthickness=0)
    logt.pack(fill="both",expand=True,padx=6,pady=6)

    def log(msg,color=None):
        logt.configure(state="normal")
        tag=f"c{color or 'def'}"
        logt.tag_configure(tag,foreground=color or C["text"])
        logt.insert("end",msg+"\n",tag); logt.see("end")
        logt.configure(state="disabled")

    def do_import():
        path=pv.get().strip()
        if not path or not os.path.exists(path):
            app._toast("Выберите файл"); return
        logt.configure(state="normal"); logt.delete("1.0","end"); logt.configure(state="disabled")
        try:
            wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
            ws=wb.active; rows=iter(ws.iter_rows(values_only=True))
            next(rows)  # заголовок
            added=skipped=errors=0
            for rn,row in enumerate(rows,2):
                if not any(row): continue
                try:
                    def cell(index, default=""):
                        return row[index] if index < len(row) else default

                    qid=f"EMP-{datetime.now().strftime('%Y%m%d%H%M%S%f')}-{rn}"
                    data = {
                        "qr_code": qid,
                        "district": cell(0),
                        "unit": cell(1),
                        "rank": cell(2),
                        "last_name": cell(3),
                        "first_name": cell(4),
                        "middle_name": cell(5),
                        "phone": cell(6),
                        "issued_date": cell(7) or datetime.now().strftime("%Y-%m-%d"),
                        "days_count": cell(8) if cell(8) not in (None, "") else 30,
                    }
                    try:
                        data = validate_pass_data(data, app.db)
                    except ValidationError as ex:
                        log(f"  Строка {rn}: пропущена ({'; '.join(ex.errors)})",C["yellow"])
                        skipped+=1
                        continue

                    full=f"{data['last_name']} {data['first_name']} {data['middle_name']}".strip()
                    inserted = insert_pass_ignore(app.db, data)
                    dl="месяц" if data["days_count"]==30 else f"{data['days_count']}сут"
                    if inserted:
                        try:
                            save_qr_code(data["qr_code"])
                        except Exception:
                            delete_pass_forever(app.db, data["qr_code"])
                            raise
                        safe_record_action(app.db, app.user, "pass.import_create", "pass", data["qr_code"], {
                            "name": full,
                            "source": os.path.basename(path),
                        })
                        log(f"  ✓  {full}  [{data['district']} / {data['unit']}]  {dl}",C["green"]); added+=1
                    else:
                        log(f"  Строка {rn}: пропущена (дубликат QR)",C["yellow"]); skipped+=1
                except Exception as ex:
                    logger.exception("Failed to import Excel row %s", rn)
                    log(f"  ✗  Строка {rn}: {ex}",C["red"]); errors+=1

            log(f"\n─── Итого: добавлено {added}, пропущено {skipped}, ошибок {errors} ───",C["accent"])
            safe_record_action(app.db, app.user, "import.excel", "file", os.path.basename(path), {
                "added": added,
                "skipped": skipped,
                "errors": errors,
            })
            if added: app._toast(f"Импортировано {added} записей",C["green"])
        except Exception as ex:
            logger.exception("Excel import failed for file: %s", path)
            log(f"Ошибка файла: {ex}",C["red"])

    def export_for_kpp():
        initial = f"kpp_passes_{datetime.now().strftime('%Y%m%d_%H%M')}{PACKAGE_EXTENSION}"
        path = filedialog.asksaveasfilename(
            parent=app,
            title="Сохранить пакет для КПП",
            initialdir=str(_desktop_dir()),
            initialfile=initial,
            defaultextension=PACKAGE_EXTENSION,
            filetypes=[("Пакет КПП", f"*{PACKAGE_EXTENSION}")],
        )
        if not path:
            return
        try:
            user_name = (app.user or {}).get("username") or (app.user or {}).get("name") or ""
            result = export_kpp_package(app.db, path, created_by=user_name)
            safe_record_action(app.db, app.user, "kpp.export", "package", result["path"].name, {
                "count": result["count"],
                "counts": result["counts"],
            })
            log(f"✓ Пакет для КПП создан: {result['path']}", C["green"])
            log(_package_summary(result), C["accent"])
            app._toast("Пакет КПП сохранён", C["green"])
        except Exception as ex:
            logger.exception("Failed to export KPP package")
            messagebox.showerror("Пакет КПП", f"Не удалось создать пакет:\n{ex}")

    def import_from_kpp_package():
        path = filedialog.askopenfilename(
            parent=app,
            title="Выбрать пакет КПП",
            filetypes=[("Пакет КПП", f"*{PACKAGE_EXTENSION}"), ("Все файлы", "*.*")],
        )
        if not path:
            return
        try:
            info = inspect_kpp_package(path)
            if not messagebox.askyesno(
                "Импорт пакета КПП",
                f"Импортировать пакет?\n\n{_package_summary(info)}\n\nПеред импортом будет создан бэкап базы.",
            ):
                return
            result = import_kpp_package(app.db, path, create_backup_before=True)
            stats = result["stats"]
            safe_record_action(app.db, app.user, "kpp.import", "package", os.path.basename(path), {
                "created": stats["created"],
                "updated": stats["updated"],
                "errors": stats["errors"],
            })
            log(f"✓ Пакет импортирован: {os.path.basename(path)}", C["green"])
            log(f"Добавлено: {stats['created']}, обновлено: {stats['updated']}, ошибок: {stats['errors']}", C["accent"])
            for error in result["errors"][:20]:
                log(f"  ✗ {error}", C["red"])
            if len(result["errors"]) > 20:
                log(f"  ... ещё ошибок: {len(result['errors']) - 20}", C["red"])
            app._toast("Пакет КПП импортирован", C["green"] if not stats["errors"] else C["yellow"])
        except Exception as ex:
            logger.exception("Failed to import KPP package")
            messagebox.showerror("Пакет КПП", f"Не удалось импортировать пакет:\n{ex}")

    Btn(win,text="Начать импорт Excel",cmd=do_import,variant="primary",
        w=612,h=44,fs=13,bg=C["panel"]).pack(padx=28,pady=(10,6))
    package_buttons=tk.Frame(win,bg=C["panel"])
    package_buttons.pack(fill="x",padx=28,pady=(0,12))
    Btn(package_buttons,text="Экспорт для КПП",cmd=export_for_kpp,variant="success",
        w=294,h=42,fs=12,bg=C["panel"]).pack(side="left")
    Btn(package_buttons,text="Импорт пакета КПП",cmd=import_from_kpp_package,variant="ghost",
        w=294,h=42,fs=12,bg=C["panel"]).pack(side="right")
