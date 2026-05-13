import os
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

from app.core.config import C, FONT as F
from app.core.logging import get_logger
from app.services.pass_db import logs_for_day, stats_summary
from app.ui.widgets import Btn

logger = get_logger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    HAS_XL = True
except ImportError:
    HAS_XL = False

def show_stats(app):
    app._clr(app.content)
    app._pgtitle.configure(text="Статистика")
    wrap=tk.Frame(app.content,bg=C["bg"]); wrap.pack(fill="both",expand=True,padx=36,pady=20)

    today=datetime.now().strftime("%Y-%m-%d")
    summary = stats_summary(app.db, today)
    total = summary["total"]; act = summary["active"]; inact = summary["inactive"]
    gd = summary["granted_today"]; dd = summary["denied_today"]; ga = summary["granted_total"]

    hdr=tk.Frame(wrap,bg=C["bg"]); hdr.pack(fill="x",pady=(0,16))
    tk.Label(hdr,text="Статистика",bg=C["bg"],fg=C["text"],font=(F,16,"bold")).pack(side="left")
    Btn(hdr,text="Экспорт дня в Excel",cmd=app._export,variant="success",
        w=210,h=38,fs=11,bg=C["bg"]).pack(side="right")

    cf=tk.Frame(wrap,bg=C["bg"]); cf.pack(fill="x")
    for i,(title,val,color) in enumerate([
        ("Всего пропусков",  total, C["accent"]),
        ("Активных",         act,   C["green"]),
        ("Просрочено",       inact, C["red"]),
        ("Проходов сегодня", gd,    C["green"]),
        ("Отказов сегодня",  dd,    C["red"]),
        ("Проходов всего",   ga,    C["accent"]),
    ]):
        card=tk.Frame(cf,bg=C["panel"],highlightthickness=1,highlightbackground=C["border"])
        card.grid(row=i//3,column=i%3,padx=8,pady=8,sticky="nsew")
        cf.columnconfigure(i%3,weight=1)
        tk.Label(card,text=str(val),bg=C["panel"],fg=color,font=(F,36,"bold")).pack(pady=(20,4))
        tk.Label(card,text=title,bg=C["panel"],fg=C["muted"],font=(F,11)).pack(pady=(0,20))

    tk.Label(wrap,text=f"События за {today}",bg=C["bg"],fg=C["muted"],font=(F,11)).pack(anchor="w",pady=(20,6))
    lt=ttk.Treeview(wrap,style="T.Treeview",columns=("t","n","q","r"),show="headings",height=9)
    for col,txt,w in [("t","Время",120),("n","ФИО",220),("q","QR",180),("r","Результат",100)]:
        lt.heading(col,text=txt); lt.column(col,width=w)
    lt.tag_configure("ok",foreground=C["green"]); lt.tag_configure("bad",foreground=C["red"])
    lt.pack(fill="x")
    for row in logs_for_day(app.db, today, descending=True):
        try:
            ts=datetime.fromisoformat(row[0]).strftime("%H:%M:%S")
        except (TypeError, ValueError):
            logger.warning("Invalid log timestamp in stats: %r", row[0])
            ts=row[0]
        lt.insert("","end",tags=("ok" if row[3] else "bad",),
                  values=(ts,row[1],row[2],"Разрешён" if row[3] else "Отказано"))

# ─────────────────────────────────────────
#  ЭКСПОРТ В EXCEL
# ─────────────────────────────────────────

def _export(app):
    if not HAS_XL: messagebox.showerror("Ошибка","pip install openpyxl"); return
    today=datetime.now().strftime("%Y-%m-%d")
    rows=logs_for_day(app.db, today, descending=False)
    if not rows: app._toast("Нет данных за сегодня",C["yellow"]); return
    path=filedialog.asksaveasfilename(defaultextension=".xlsx",
         initialfile=f"статистика_{today}.xlsx",filetypes=[("Excel","*.xlsx")])
    if not path: return
    wb=openpyxl.Workbook(); ws=wb.active; ws.title=f"Статистика {today}"
    hf=PatternFill("solid",fgColor="1E2535")
    okf=PatternFill("solid",fgColor="162B20"); bdf=PatternFill("solid",fgColor="2B1620")
    thin=Side(style="thin",color="2E3A52")
    brd=Border(left=thin,right=thin,top=thin,bottom=thin)
    ctr=Alignment(horizontal="center",vertical="center")
    ws.merge_cells("A1:D1"); ws["A1"]=f"Статистика — {today}"
    ws["A1"].font=Font(bold=True,size=14,color="4F8EF7"); ws["A1"].alignment=ctr
    ws.row_dimensions[1].height=30
    total=len(rows); gr=sum(1 for r in rows if r[3]); dn=total-gr
    ws.merge_cells("A2:D2"); ws["A2"]=f"Всего: {total}  |  Разрешено: {gr}  |  Отказано: {dn}"
    ws["A2"].font=Font(italic=True,color="8A9BBF",size=10); ws["A2"].alignment=ctr
    ws.row_dimensions[3].height=6
    for ci,h in enumerate(["Время","ФИО","QR-код","Результат"],1):
        cell=ws.cell(4,ci,h); cell.fill=hf
        cell.font=Font(bold=True,color="E2E8F5",size=11)
        cell.alignment=ctr; cell.border=brd
    ws.row_dimensions[4].height=22
    for ri,(ts,name,qr,ok) in enumerate(rows,5):
        try:
            t=datetime.fromisoformat(ts).strftime("%H:%M:%S")
        except (TypeError, ValueError):
            logger.warning("Invalid log timestamp during export: %r", ts)
            t=ts
        for ci,val in enumerate([t,name,qr,"Разрешён" if ok else "Отказано"],1):
            cell=ws.cell(ri,ci,val)
            cell.fill=okf if ok else bdf
            cell.font=Font(color=("2EC27E" if ok else "E05567") if ci==4 else "E2E8F5",size=11)
            cell.alignment=ctr if ci in (1,4) else Alignment(vertical="center")
            cell.border=brd
        ws.row_dimensions[ri].height=20
    for col,w in [("A",14),("B",30),("C",26),("D",14)]:
        ws.column_dimensions[col].width=w
    ws.freeze_panes="A5"
    try: wb.save(path); app._toast(f"Сохранено: {os.path.basename(path)}",C["green"])
    except Exception as ex:
        logger.exception("Failed to save Excel export: %s", path)
        messagebox.showerror("Ошибка",str(ex))

# ─────────────────────────────────────────
#  ПОЛЬЗОВАТЕЛИ
# ─────────────────────────────────────────
